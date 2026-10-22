"""HS material report: every material code with its activity code (e.g. PPCFD), activity name, and sales split
into Retail and Project, with totals. Reads the extract inside 'Home Solutions Dashboard.twbx' (or an HS
dashboard built from it), so it shows the same data as the dashboard after its last refresh.

usage:
    python material_report.py "Home Solutions Dashboard.twbx" "HS material report.xlsx"
        [--start 2026-01-01] [--end 2026-12-31] [--measure gross|net] [--units INR|Lakh|Mn|Cr] [--without-fd]

Rules (same as the HS dashboard):
  * Sales = rows of type 'Sales', dated by invoice date, between --start and --end (default: this year to date).
  * Activity: HDCWD / HSCWD = WD, HDCFD / PPCFD = FD, HDCOS / HSCOS = COS, HDSPE / HSSPE = SPE, PREVC = EVC.
  * GTM: Retail when the pricing class is EKR or ERW, or the activity is FD; otherwise Project.
  * Net = gross x (1 - rebate): Retail 13%, Project 5.1% for this year's sales; 12.2% / 4.5% for earlier years.
    (--measure net-source uses the Net Sales column of the source instead.)

needs: pip install tableauhyperapi pandas openpyxl
"""
import argparse
import datetime as dt
import os
import shutil
import tempfile
import zipfile

import pandas as pd
from tableauhyperapi import Connection, HyperProcess, Telemetry

ACTIVITY = {"HDCWD": "WD", "HSCWD": "WD", "HDCFD": "FD", "PPCFD": "FD", "HDCOS": "COS", "HSCOS": "COS",
            "HDSPE": "SPE", "HSSPE": "SPE", "PREVC": "EVC"}
ACTIVITY_NAME = {"WD": "Wiring Devices", "FD": "Final Distribution", "COS": "COS", "SPE": "SPE", "EVC": "EV Charging",
                 "Other": "Other"}
REBATE = {("Retail", "CY"): 0.13, ("Project", "CY"): 0.051, ("Retail", "PY"): 0.122, ("Project", "PY"): 0.045}
UNITS = {"INR": 1, "Lakh": 1e5, "Mn": 1e6, "Cr": 1e7}


def read_sales(twbx, start, end):
    """Sales rows of the extract between start and end (inclusive)."""
    tmp = tempfile.mkdtemp()
    try:
        with zipfile.ZipFile(twbx) as z:
            hyper = [n for n in z.namelist() if n.endswith(".hyper") and "federated_14wc2mc1wc09dr1b4kqf30" in n]
            if not hyper:
                raise SystemExit("no HS extract (federated_14wc2mc1wc09dr1b4kqf30*.hyper) in " + twbx)
            path = z.extract(hyper[0], tmp)
        with HyperProcess(telemetry=Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU, parameters={"log_dir": tmp}) as hp, \
                Connection(hp.endpoint, path) as c:
            table = [t for t in c.catalog.get_table_names("Extract") if t.name.unescaped.startswith("Year_")][0]
            cols = {str(x.name).strip('"') for x in c.catalog.get_table_definition(table).columns}
            fam = '"Family Actual"' if "Family Actual" in cols else "NULL"
            q = (f'SELECT "MaterialCode", "MaterialDescription", {fam} AS "Family", "HeliosCode", "PricingClassL3", '
                 f'EXTRACT(YEAR FROM "InvoiceDate") AS "Year", "Gross Sales", "Net Sales", "BillingQuantity" FROM {table} '
                 f"WHERE \"Type\" = 'Sales' AND \"InvoiceDate\" >= '{start}' AND \"InvoiceDate\" < '{end + dt.timedelta(days=1)}'")
            rows = c.execute_list_query(q)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    df = pd.DataFrame(rows, columns=["Material code", "Material description", "Family", "Activity code", "Pricing class",
                                     "Year", "Gross", "Net (source)", "Quantity"])
    for col in ("Gross", "Net (source)"):
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    df["Quantity"] = pd.to_numeric(df["Quantity"], errors="coerce").fillna(0).astype(int)
    return df


def build_report(df, measure, units, current_year, without_fd):
    df = df.copy()
    df["Activity"] = df["Activity code"].map(ACTIVITY).fillna("Other")
    df["Activity name"] = df["Activity"].map(ACTIVITY_NAME)
    retail = df["Pricing class"].isin(["EKR", "ERW"]) | (df["Activity"] == "FD")
    df["GTM"] = retail.map({True: "Retail", False: "Project"})
    if without_fd:
        df = df[df["Activity"] != "FD"]
    year = df["Year"].astype(int)
    if measure == "gross":
        value = df["Gross"]
    elif measure == "net-source":
        value = df["Net (source)"]
    else:
        rate = [REBATE[(g, "CY" if y >= current_year else "PY")] for g, y in zip(df["GTM"], year)]
        value = df["Gross"] * (1 - pd.Series(rate, index=df.index))
    df["Value"] = value / UNITS[units]

    keys = ["Material code", "Material description", "Family", "Activity code", "Activity", "Activity name"]
    df[keys] = df[keys].fillna("Unmapped")
    sales = df.pivot_table(index=keys, columns="GTM", values="Value", aggfunc="sum", fill_value=0.0)
    qty = df.pivot_table(index=keys, columns="GTM", values="Quantity", aggfunc="sum", fill_value=0)
    out = pd.DataFrame(index=sales.index)
    for g in ("Retail", "Project"):
        out[f"{g} sales"] = sales[g] if g in sales else 0.0
    out["Total sales"] = out["Retail sales"] + out["Project sales"]
    out["Retail share %"] = (out["Retail sales"] / out["Total sales"].where(out["Total sales"] != 0)).fillna(0.0)
    for g in ("Retail", "Project"):
        out[f"{g} quantity"] = qty[g] if g in qty else 0
    out["Total quantity"] = out["Retail quantity"] + out["Project quantity"]
    out = out.reset_index().sort_values(["Activity", "Total sales"], ascending=[True, False])

    by_act = out.groupby(["Activity code", "Activity", "Activity name"], as_index=False)[
        ["Retail sales", "Project sales", "Total sales", "Retail quantity", "Project quantity", "Total quantity"]].sum()
    by_act["Retail share %"] = (by_act["Retail sales"] / by_act["Total sales"].where(by_act["Total sales"] != 0)).fillna(0.0)
    by_act["Materials"] = out.groupby("Activity code")["Material code"].nunique().reindex(by_act["Activity code"]).values
    return add_total(out, "Material code"), add_total(by_act, "Activity code")


def add_total(df, label_col):
    num = df.select_dtypes("number").columns
    total = {c: df[c].sum() for c in num}
    if "Retail share %" in df:
        total["Retail share %"] = total["Retail sales"] / total["Total sales"] if total["Total sales"] else 0.0
    total[label_col] = "Grand Total"
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)


def write_excel(path, detail, by_act, info):
    from openpyxl.styles import Font, PatternFill
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(info, columns=["Setting", "Value"]).to_excel(xw, sheet_name="About", index=False)
        by_act.to_excel(xw, sheet_name="By activity", index=False)
        detail.to_excel(xw, sheet_name="Materials", index=False)
        head = PatternFill("solid", fgColor="1F7A4D")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
            for col in ws.columns:
                name = str(col[0].value)
                ws.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or "")) for c in col) + 2))
                fmt = "0.0%" if name.endswith("%") else ("#,##0" if "quantity" in name.lower() or name == "Materials" else "#,##0.0")
                for c in col[1:]:
                    if isinstance(c.value, (int, float)):
                        c.number_format = fmt
            last = ws.max_row
            if ws.title != "About" and ws.cell(last, 1).value == "Grand Total":
                for c in ws[last]:
                    c.font = Font(bold=True)
            ws.freeze_panes = "B2"


def main():
    today = dt.date.today()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("twbx")
    ap.add_argument("out")
    ap.add_argument("--start", type=dt.date.fromisoformat, default=dt.date(today.year, 1, 1))
    ap.add_argument("--end", type=dt.date.fromisoformat, default=today)
    ap.add_argument("--measure", choices=["gross", "net", "net-source"], default="gross")
    ap.add_argument("--units", choices=list(UNITS), default="Mn")
    ap.add_argument("--without-fd", action="store_true", help="leave out FD (HDCFD / PPCFD)")
    a = ap.parse_args()

    df = read_sales(a.twbx, a.start, a.end)
    detail, by_act = build_report(df, a.measure, a.units, a.end.year, a.without_fd)
    info = [("Source", os.path.basename(a.twbx)), ("Period", f"{a.start} to {a.end}"),
            ("Measure", a.measure + (" (gross x (1 - rebate): Retail 13% / Project 5.1% this year)" if a.measure == "net" else "")),
            ("Units", a.units + " (sales); quantity in units sold"), ("FD", "excluded" if a.without_fd else "included"),
            ("GTM", "Retail = pricing class EKR / ERW, or activity FD; the rest is Project"),
            ("Materials", int(detail["Material code"].ne("Grand Total").sum())),
            ("Created", dt.datetime.now().strftime("%Y-%m-%d %H:%M"))]
    write_excel(a.out, detail, by_act, info)
    tot = detail.iloc[-1]
    print(f"{a.out}: {info[6][1]} materials; Retail {tot['Retail sales']:,.1f} + Project {tot['Project sales']:,.1f} "
          f"= {tot['Total sales']:,.1f} {a.units}")


if __name__ == "__main__":
    main()
