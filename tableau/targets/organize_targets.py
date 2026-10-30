"""Organise the Retail target files (Master Targets File + QTOD and ATOD file) and the HS Project targets into one
long-format targets workbook for the Retail, HS and Wiser dashboards, splitting each target by last year's sales
contribution.

The same code runs as the last chunk of Channel_Analytics_v9.ipynb (section 10), where the sales come from
df_Sales_final. Run on its own (tests, one-off use):

    python organize_targets.py MasterTargets.xlsx QTOD.xlsx sales.parquet out_dir [--hs HS_Targets.xlsx] [--year 2026]

`sales` is one row per sale (or pre-aggregated) with the columns of SALES_COLUMNS.

How targets are organised (base year = target year - 1, gross sales):
  * Yearly and monthly targets are separate, as in the Retail file: 'Branch Targets' (yearly) and 'Monthly Targets'
    (the monthly commit, which need not add up to the yearly target). Both are kept exactly as given; months that
    have no target yet stay empty — nothing is derived from the yearly target.
  * Product groups: each yearly or monthly target is split over product groups by the branch's base-year sales mix.
    Monthly Miluz Zeta and Wiser targets are kept as given and only the rest is split. Activities (WD, FD, COS, SPE,
    EVC) are the sum of their product groups.
  * Project: yearly targets (HS 'Branch Targets', by activity) and monthly commits (HS 'Activity wise Month Target')
    split over product groups within the activity. The HS file's Retail rows are replaced by the Retail file, so the
    Retail and HS dashboards show the same Retail targets.
  * Partners: the quarter targets are kept; Zeta / Wiser partner targets go to quarters like the partner's total.
    'Partner Product Split' splits each partner's yearly target over product groups by its own base-year mix (the
    branch mix for partners with little or no base-year sales).
Every split row says where it came from (column 'Source'), so given and derived numbers are never confused.
"""
import argparse
import datetime as dt
import re

import numpy as np
import pandas as pd

SALES_COLUMNS = ["Year", "Month", "GTM", "Branch", "Region", "HeliosCode", "Family", "Partner", "Gross"]

ACTIVITY = {"HDCWD": "WD", "HSCWD": "WD", "HDCFD": "FD", "PPCFD": "FD", "HDFDA": "FD", "HDFDB": "FD", "HDFDS": "FD",
            "HDCOS": "COS", "HSCOS": "COS", "HDSPE": "SPE", "HSSPE": "SPE", "PREVC": "EVC"}
ACTIVITIES = ["WD", "FD", "COS", "SPE", "EVC"]
# product groups as in the 'Product Targets' sheet, with the activity each belongs to
GROUPS = [("Unica Pure", "WD"), ("Livia", "WD"), ("AvatarOn", "WD"), ("Miluz Zeta", "WD"), ("WD- All others", "WD"),
          ("Easy9 6kA", "FD"), ("Easy9 10kA", "FD"), ("Easy9 RCCB", "FD"), ("Easy9 DB", "FD"), ("FD- All others", "FD"),
          ("Wiser", "COS"), ("COS- All others", "COS"), ("SPE", "SPE"), ("EVC", "EVC")]
GROUP_ACTIVITY = dict(GROUPS)
FOCUS = {"Miluz Zeta": "Miluz Zeta", "Wiser": "Wiser"}      # monthly target column -> product group
FOCUS_FAMILY = {"Zeta Target": "Miluz Zeta", "Wiser Target": "Wiser"}   # QTOD columns

# branch spellings in the target files -> the branch the sales data uses
BRANCH_ALIASES = {"NCR(DEL+ROH)": "NCR", "NCR(NOI+GHAZ)": "NOIDA-GHZ", "PNQ": "PUNE", "NAN": "NASHIK", "TN2": "TN-2",
                  "TN -1": "TN-1", "TN-1 ": "TN-1", "AP ": "AP", "ap": "AP", "UP": "UP-1", "MP": "INDORE"}
SSP_ALIASES = {"Katson Agency": "Kaston Agency", "Balaji Sales": "Balaji Sales Corp"}
QUARTERS = {"Q1": (1, 2, 3), "Q2": (4, 5, 6), "Q3": (7, 8, 9), "Q4": (10, 11, 12)}
MIN_PARTNER_SALES = 1e5          # partners below this base-year sales get their branch's product mix


def product_group(activity, family):
    f = str(family or "")
    if activity == "WD":
        if f in ("Unica Pure", "Unica Pure Plus"):
            return "Unica Pure"
        return f if f in ("Livia", "AvatarOn", "Miluz Zeta") else "WD- All others"
    if activity == "FD":
        if f.startswith("Easy9 MCB 6kA"):
            return "Easy9 6kA"
        if f.startswith("Easy9 MCB 10kA"):
            return "Easy9 10kA"
        if f == "Easy9 RCCB":
            return "Easy9 RCCB"
        if f.startswith("Easy9 DB") or f in ("Easy9 Enclosure", "Easy9 Metal P and S"):
            return "Easy9 DB"
        return "FD- All others"
    if activity == "COS":
        return "Wiser" if f == "Wiser" else "COS- All others"
    return activity if activity in ("SPE", "EVC") else None


def canon_branch(b):
    if b is None or (isinstance(b, float) and np.isnan(b)) or str(b).strip() in ("", "0", "nan"):
        return None
    s = str(b).strip()
    return BRANCH_ALIASES.get(s, BRANCH_ALIASES.get(s.upper(), s))


def clean_code(v):
    return re.sub(r"\.0$", "", str(v).strip()) if pd.notna(v) else ""


def num(s):
    return pd.to_numeric(s, errors="coerce")


# ====================================================================== sales basis
def prepare_sales(sales):
    s = sales.copy()
    s["Activity"] = s["HeliosCode"].astype(str).str.strip().str.upper().map(ACTIVITY)
    s["Group"] = [product_group(a, f) for a, f in zip(s["Activity"], s["Family"])]
    retail = s["GTM"] == "Retail"          # the aliases are Retail spellings; Project uses NS Branches names (NAN, MP ...)
    s.loc[retail, "Branch"] = s.loc[retail, "Branch"].map(canon_branch)
    s["Partner"] = s["Partner"].map(clean_code)
    s["Gross"] = num(s["Gross"]).fillna(0.0)
    return s[s["Group"].notna()]


def shares(df, keys, col, weight="Gross"):
    """Share of `col` within `keys` (e.g. month within branch)."""
    g = df.groupby(keys + [col])[weight].sum()
    tot = g.groupby(level=list(range(len(keys)))).transform("sum")
    return (g / tot.where(tot > 0)).fillna(0.0)


def lookup_share(table, key, values, fallback):
    """{value: share} for one key, or the fallback table's shares when the key has no sales."""
    try:
        s = table.loc[key]
        if s.sum() > 0:
            return s.reindex(values, fill_value=0.0)
    except KeyError:
        pass
    return fallback.reindex(values, fill_value=0.0)


def contribution(sales, base_year):
    """Base-year sales contribution tables (values in rupees and shares)."""
    s = sales[sales["Year"] == base_year]
    out = {}

    def table(part, keys, col, name):
        v = part.groupby(keys + [col])["Gross"].sum().unstack(col, fill_value=0.0)
        if col == "Month":
            v.columns = [dt.date(2000, int(m), 1).strftime("%b") for m in v.columns]
        tot = v.sum(axis=1)
        sh = v.div(tot.where(tot > 0), axis=0).fillna(0.0)
        sh.columns = [f"{c} %" for c in sh.columns]
        out[name] = pd.concat([v.assign(Total=tot), sh], axis=1).reset_index()

    for g in ("Retail", "Project"):
        part = s[s["GTM"] == g]
        if part.empty:
            continue
        table(part, ["Region", "Branch"], "Month", f"{g} branch x month")
        table(part, ["Region", "Branch"], "Activity", f"{g} branch x activity")
        table(part, ["Region", "Branch"], "Group", f"{g} branch x product")
        table(part, ["Region"], "Group", f"{g} region x product")
    v = s.groupby(["Region", "Branch", "GTM"])["Gross"].sum().unstack("GTM", fill_value=0.0)
    tot = v.sum(axis=1)
    out["Retail vs Project"] = v.assign(Total=tot, **{f"{c} %": v[c] / tot.where(tot > 0) for c in v.columns}).reset_index()
    p = s[s["GTM"] == "Retail"].groupby(["Branch", "Partner", "Group"])["Gross"].sum().unstack("Group", fill_value=0.0)
    out["Retail partner x product"] = p.assign(Total=p.sum(axis=1)).sort_values("Total", ascending=False).reset_index()
    return out


# ====================================================================== organise
class Issues(list):
    def add(self, sheet, what, action, rows=""):
        self.append({"Source sheet": sheet, "Problem": what, "What was done": action, "Rows / values": str(rows)[:500]})


def _mix_rows(total, mix, focus, keys, source_split):
    """Rows of one target split over product groups: `focus` groups as given, the rest by `mix`."""
    rest = total - sum(v for v, _ in focus.values())
    others = mix.drop(list(focus), errors="ignore")
    others = others / others.sum() if others.sum() > 0 else others
    rows = []
    for grp, act in GROUPS:
        if grp in focus:
            v, how = focus[grp]
        else:
            v, how = max(rest, 0.0) * others.get(grp, 0.0), source_split
        if v:
            rows.append((*keys, act, grp, v, "Net", how))
    return rows, rest


def retail_targets(master, sales, year, issues):
    """Retail yearly targets ('Branch Targets') and monthly targets ('Monthly Targets' = the monthly commit), each as
    given and split over product groups by the branch's base-year mix. The two are independent: months are never
    derived from the yearly target and need not add up to it."""
    base = sales[(sales["Year"] == year - 1) & (sales["GTM"] == "Retail")]
    grp_sh = shares(base, ["Branch"], "Group")
    reg_sh = shares(base, ["Region"], "Group")
    grp_all = base.groupby("Group")["Gross"].sum() / base["Gross"].sum()
    names = [g for g, _ in GROUPS]

    bt = master["Branch Targets"].copy()
    bt["Branch"] = bt["Branch"].map(canon_branch)
    region = dict(zip(bt["Branch"], bt["Zone"]))
    yearly = bt[bt["Year"] == year].groupby("Branch")["Target"].sum(min_count=1).dropna()
    if bt.loc[bt["Year"] == year - 1, "Target"].isna().all() and (bt["Year"] == year - 1).any():
        issues.add("Branch Targets", f"{year - 1} rows have no target", "ignored")

    mt = master["Monthly Targets"].iloc[:, :6].copy()
    mt.columns = ["Zone", "Branch", "Month", "Target", "Miluz Zeta", "Wiser"]
    renamed = sorted({b for b in mt["Branch"].dropna().astype(str).str.strip() if canon_branch(b) != b})
    if renamed:
        issues.add("Monthly Targets", "branch names that differ from the sales data",
                   "renamed: " + ", ".join(f"{b} -> {canon_branch(b)}" for b in renamed))
    side = master["Monthly Targets"].iloc[:, 6:].dropna(how="all")
    if not side.empty:
        issues.add("Monthly Targets", "a second table to the right of the data (e.g. 'Oct WD / FD / Wiser Plan', in lakh)",
                   "not read; keep only the data columns in this sheet", f"{len(side)} rows")
    mt["Branch"] = mt["Branch"].map(canon_branch)
    mt["Month"] = pd.to_datetime(mt["Month"], errors="coerce")
    mt = mt[mt["Month"].dt.year == year]
    for c in ("Target", "Miluz Zeta", "Wiser"):
        mt[c] = num(mt[c])
    mt["m"] = mt["Month"].dt.month
    monthly = mt.groupby(["Branch", "m"])[["Target", "Miluz Zeta", "Wiser"]].sum(min_count=1).dropna(subset=["Target"])
    have = sorted(monthly.index.get_level_values(1).unique())
    if len(have) < 12:
        missing = [dt.date(year, m, 1).strftime("%b") for m in range(1, 13) if m not in have]
        issues.add("Monthly Targets", f"no monthly targets yet for {', '.join(missing)}",
                   "left empty (monthly targets are never derived from the yearly target) — add them when set")
    for c in ("Miluz Zeta", "Wiser"):
        blank = mt[mt["Target"].notna() & mt[c].isna()]
        if len(blank):
            issues.add("Monthly Targets", f"blank {c} target", "taken as 0",
                       ", ".join(f"{b} {d:%b}" for b, d in zip(blank["Branch"], blank["Month"])))
    for b in sorted(set(monthly.index.get_level_values(0)) - set(region)):
        issues.add("Monthly Targets", f"branch {b} is not in 'Branch Targets'", "region left blank")

    def mix_of(b):
        return lookup_share(grp_sh, b, names, lookup_share(reg_sh, region.get(b), names, grp_all))

    split = "Split by last year's product mix"
    y_rows, m_rows = [], []
    for b, y in yearly.items():
        rows, _ = _mix_rows(y, mix_of(b), {}, (year, region.get(b), b, "Retail"), "Yearly as given; " + split.lower())
        y_rows += rows
    for (b, m), r in monthly.iterrows():
        focus = {grp: (float(r[col]) if pd.notna(r[col]) else 0.0, "Given") for col, grp in FOCUS.items()}
        rows, rest = _mix_rows(r["Target"], mix_of(b), focus, (year, dt.date(year, m, 1), region.get(b), b, "Retail"),
                               "Given month; " + split.lower())
        if rest < 0:
            issues.add("Monthly Targets", f"{b} {dt.date(year, m, 1):%b}: Zeta + Wiser above the branch target",
                       "other product groups 0")
        m_rows += rows
    tm = monthly["Target"].groupby(level=0).sum()
    over = [(b, tm[b], yearly[b]) for b in tm.index if b in yearly and tm[b] > yearly[b]]
    if over:
        issues.add("Monthly Targets", "monthly targets already add up to more than the yearly target",
                   "both kept as given (they are separate targets) — just for information",
                   "; ".join(f"{b}: months {t / 1e7:,.2f} Cr vs yearly {y / 1e7:,.2f} Cr" for b, t, y in over))
    ycols = ["Year", "Region", "Branch", "GTM", "Activity", "Product group", "Target", "Basis", "Source"]
    mcols = ["Year", "Month", "Region", "Branch", "GTM", "Activity", "Product group", "Target", "Basis", "Source"]
    return pd.DataFrame(y_rows, columns=ycols), pd.DataFrame(m_rows, columns=mcols), yearly, tm, region


def _project_mix(base, b, act):
    sub = base[base["Activity"] == act]
    grp = sub.groupby(["Branch", "Group"])["Gross"].sum()
    mix = grp.loc[b] if b in grp.index.get_level_values(0) and grp.loc[b].sum() > 0 else sub.groupby("Group")["Gross"].sum()
    if mix.sum() <= 0:
        mix = pd.Series({[g for g, a in GROUPS if a == act][-1]: 1.0})
    return mix / mix.sum()


def project_targets(hs, sales, year, issues):
    """Project yearly targets from the HS file's 'Branch Targets' (by activity), split over product groups."""
    if hs is None or "Branch Targets" not in hs:
        return pd.DataFrame()
    bt = hs["Branch Targets"].copy()
    rr = bt[(bt["Year"] == year) & (bt["Category"] == "Retail")]
    if len(rr) and rr["Target"].nunique() == 1:
        issues.add("HS Branch Targets", f"every Retail branch has the same {year} target ({rr['Target'].iloc[0]:,.0f})",
                   "replaced by the Retail 'Branch Targets' (Retail and HS now share one Retail target)")
    bt = bt[(bt["Year"] == year) & (bt["Category"] == "Project")]
    dup = bt[bt.duplicated(["Branch", "Helios Code"], keep=False)]
    if len(dup):
        issues.add("HS Branch Targets", "two rows for the same branch and Helios code (HDCOS twice per branch)",
                   "added together as COS", f"{dup['Branch'].nunique()} branches")
    bt["Activity"] = bt["Helios Code"].map(ACTIVITY)
    base = sales[(sales["Year"] == year - 1) & (sales["GTM"] == "Project")]
    region = base.groupby("Branch")["Region"].agg(lambda s: s.mode().iloc[0] if len(s.mode()) else None).to_dict()
    rows = []
    for (b, act), y in bt.groupby(["Branch", "Activity"])["Target"].sum().items():
        for g, sh in _project_mix(base, b, act).items():
            rows.append((year, region.get(b), b, "Project", act, g, y * sh, "Net",
                         "Yearly by activity as given; split by last year's Project product mix"))
    if (~bt["Branch"].isin(region)).any():
        issues.add("HS Branch Targets", "Project branches with no Project sales last year",
                   "product split by all Project sales; region left blank", ", ".join(sorted(set(bt["Branch"]) - set(region))))
    return pd.DataFrame(rows, columns=["Year", "Region", "Branch", "GTM", "Activity", "Product group", "Target", "Basis",
                                       "Source"])


def project_monthly(com, sales, year, issues, yearly_p=None):
    """Project monthly targets = the HS file's Project commits by activity, split over product groups.
    The HS file's Retail commits are not used: Retail monthly targets come from the Retail 'Monthly Targets' sheet,
    so the Retail and HS dashboards show the same Retail numbers."""
    if com is None or com.empty:
        return pd.DataFrame()
    if (com["GTM"] == "Retail").any():
        issues.add("Activity wise Month Target", "Retail commits in the HS file differ from the Retail 'Monthly Targets'",
                   "not used — Retail monthly targets come from the Retail file for both dashboards",
                   f"{(com['GTM'] == 'Retail').sum()} rows")
    p = com[(com["GTM"] == "Project") & com["Activity"].notna()]
    fd = p[p["Activity"] == "FD"]
    if len(fd):
        issues.add("Activity wise Month Target", "Project FD commits (FD is always Retail; values up to 124 Cr a month)",
                   "not used", f"{len(fd)} rows")
    p = p[p["Activity"] != "FD"]
    if yearly_p is not None and len(yearly_p):
        # a month whose commit equals the whole yearly target is the yearly figure pasted in, not a monthly commit
        y = yearly_p.groupby(["Branch", "Activity"])["Target"].sum()
        tot = p.groupby(["Branch", "Month", "Activity"])["Commit"].transform("sum")
        yv = pd.Series([y.get((b, a), np.nan) for b, a in zip(p["Branch"], p["Activity"])], index=p.index)
        pasted = (yv > 0) & ((tot - yv).abs() <= 0.001 * yv)
        if pasted.any():
            months = sorted(p.loc[pasted, "Month"].dt.strftime("%b").unique())
            issues.add("Activity wise Month Target", f"Project commits for {', '.join(months)} equal the whole yearly target",
                       "not used (the yearly figure pasted in as a month)", f"{p.loc[pasted, 'Branch'].nunique()} branches")
            p = p[~pasted]
    if (com["Commit"].round(0) != com["OB Commit"].round(0)).mean() < 0.01:
        issues.add("Activity wise Month Target", "'OB Commit' is the same as 'Commit' on almost every row",
                   "only 'Commit' kept (as the monthly target)")
    base = sales[(sales["Year"] == year - 1) & (sales["GTM"] == "Project")]
    rows = []
    for (b, m, reg, act), v in p.groupby(["Branch", "Month", "Region", "Activity"], dropna=False)["Commit"].sum().items():
        for g, sh in _project_mix(base, b, act).items():
            rows.append((year, pd.Timestamp(m).date(), reg, b, "Project", act, g, v * sh, "Net",
                         "Given commit by activity; split by last year's Project product mix"))
    return pd.DataFrame(rows, columns=["Year", "Month", "Region", "Branch", "GTM", "Activity", "Product group", "Target",
                                       "Basis", "Source"])


def commits(hs, year, issues):
    """The HS file's 'Activity wise Month Target' rows for the year (Commit by branch, month, GTM and activity)."""
    if hs is None or "Activity wise Month Target" not in hs:
        return pd.DataFrame()
    c = hs["Activity wise Month Target"].copy()
    c["Month"] = pd.to_datetime(c["Month and year"], errors="coerce")
    c = c[c["Month"].dt.year == year]
    for k in ("Commit", "OB Commit"):
        c[k] = num(c[k])
    c["Activity"] = c["Helios Code"].map(ACTIVITY)
    months = sorted(c.loc[c["Category"] == "Project", "Month"].dt.month.unique())
    if months and len(months) < 12:
        issues.add("Activity wise Month Target", f"Project commits only for {len(months)} months "
                   f"({', '.join(dt.date(year, m, 1).strftime('%b') for m in months)})", "other months left empty")
    out = c.rename(columns={"Zone": "Region", "Category": "GTM"})
    retail = out["GTM"] == "Retail"
    out.loc[retail, "Branch"] = out.loc[retail, "Branch"].map(canon_branch)
    return out[["Month", "Region", "Branch", "GTM", "Activity", "Commit", "OB Commit"]]


def partner_targets(qtod, master, sales, year, partner_info, region, issues):
    q = qtod.copy()
    q.columns = [str(c).strip() for c in q.columns]
    q = q[q["Target Year"] == year].copy()
    q["Partner Code"] = q["CC"].map(clean_code)
    qcols = ["Qtr 1 Target", "Qtr 2 Target", "Qtr 3 Target", "Qtr 4 Target"]
    for c in qcols + ["Target 100%", "Target 108%"] + list(FOCUS_FAMILY):
        q[c] = num(q[c])
    text = [c for c in qcols if qtod[c].map(lambda v: isinstance(v, str)).any()]
    if text:
        issues.add("QTOD - ATOD", f"{', '.join(text)} stored as text", "converted to numbers")
    # partner branch: partner master, else the QTOD branch (aliases fixed)
    pt = master.get("Partner Targets")
    pm_branch = {}
    if pt is not None:
        pm_branch = {clean_code(k): canon_branch(v) for k, v in zip(pt["Sold-to Party"], pt["Branch"]) if canon_branch(v)}
    info_branch = {k: canon_branch(v.get("Branch")) for k, v in partner_info.items() if canon_branch(v.get("Branch"))}
    qb = q["Branch"].map(canon_branch)
    q["Branch"] = [info_branch.get(c) or pm_branch.get(c) or b for c, b in zip(q["Partner Code"], qb)]
    diff = (qb != q["Branch"]).sum()
    codes = sorted({str(b) for b in qtod["Branch"].dropna().unique()} - set(region))
    if codes:
        issues.add("QTOD - ATOD", "branch codes not used by the sales data / Branch Targets: " + ", ".join(codes),
                   f"branch taken from the partner master ({diff} partners changed)")
    q["Partner type"] = q["Partner Code"].map(lambda c: partner_info.get(c, {}).get("RD/KRO"))
    if q["RD/ KRO"].isna().all() if "RD/ KRO" in q else True:
        issues.add("QTOD - ATOD", "'RD/ KRO' is empty", "filled from the partner master / sales data")
    q["Region"] = q["Branch"].map(region)
    bad_q = q[(q[qcols].sum(axis=1) - q["Target 100%"]).abs() > 0.01 * q["Target 100%"]]
    if len(bad_q):
        issues.add("QTOD - ATOD", "quarter targets do not add up to 'Target 100%'",
                   "kept as given (mostly partners added during the year)",
                   "; ".join(f"{p} ({c})" for p, c in zip(bad_q["Partner"], bad_q["Partner Code"]))[:500])
    ratio = q["Target 108%"] / q["Target 100%"]
    bad_r = q[(ratio - 1.08).abs() > 0.005]
    if len(bad_r):
        issues.add("QTOD - ATOD", "'Target 108%' is not 108% of 'Target 100%'",
                   "quarter 108% targets = quarter x 1.08; check the yearly 108% figure",
                   "; ".join(f"{p}: {r:.2f}x" for p, r in zip(bad_r["Partner"], ratio[bad_r.index])))
    dn = q[q["Partner"].duplicated(keep=False)]
    if len(dn):
        issues.add("QTOD - ATOD", "same partner name on different partner codes (check for duplicates)", "kept",
                   ", ".join(sorted(dn["Partner"].unique()))[:400])
    rows = []
    base = sales[(sales["Year"] == year - 1) & (sales["GTM"] == "Retail")]
    qsh_b = shares(base.assign(Q=(base["Month"] - 1) // 3 + 1), ["Branch"], "Q")
    for _, r in q.iterrows():
        qt = r[qcols].fillna(0.0).values.astype(float)
        ok = abs(qt.sum() - r["Target 100%"]) <= 0.01 * r["Target 100%"] and qt.sum() > 0
        qs = qt / qt.sum() if ok else lookup_share(qsh_b, r["Branch"], [1, 2, 3, 4], pd.Series(0.25, index=[1, 2, 3, 4])).values
        for i, qn in enumerate(QUARTERS):
            rows.append((year, qn, r["Partner Code"], r["Partner"], r["Branch"], r["Region"], "Retail", r["Partner type"],
                         "All", qt[i], qt[i] * 1.08, "Gross", "Given (108% = quarter x 1.08)"))
            for col, fam in FOCUS_FAMILY.items():
                v = r[col] if pd.notna(r[col]) else 0.0
                rows.append((year, qn, r["Partner Code"], r["Partner"], r["Branch"], r["Region"], "Retail", r["Partner type"],
                             fam, v * qs[i], v * qs[i] * 1.08, "Gross",
                             "Yearly, by the partner's quarters" if ok else "Yearly, by the branch's quarters last year"))
    cols = ["Year", "Quarter", "Partner Code", "Partner", "Branch", "Region", "GTM", "Partner type", "Family",
            "Target 100%", "Target 108%", "Basis", "Source"]
    return pd.DataFrame(rows, columns=cols), q


def partner_product_split(q, sales, year):
    base = sales[(sales["Year"] == year - 1) & (sales["GTM"] == "Retail")]
    ps = base.groupby(["Partner", "Group"])["Gross"].sum()
    ptot = ps.groupby(level=0).sum()
    bsh = shares(base, ["Branch"], "Group")
    allsh = base.groupby("Group")["Gross"].sum() / base["Gross"].sum()
    rows = []
    for _, r in q.iterrows():
        t = r["Target 100%"]
        focus = {fam: (r[col] if pd.notna(r[col]) else 0.0) for col, fam in FOCUS_FAMILY.items()}
        code = r["Partner Code"]
        if ptot.get(code, 0) >= MIN_PARTNER_SALES:
            mix, how = ps.loc[code], "Partner's product mix last year"
        else:
            mix, how = lookup_share(bsh, r["Branch"], [g for g, _ in GROUPS], allsh), "Branch's product mix last year (partner below 1 L)"
        mix = mix.drop(list(focus), errors="ignore")
        mix = mix / mix.sum() if mix.sum() > 0 else allsh.drop(list(focus)) / allsh.drop(list(focus)).sum()
        rest = max(t - sum(focus.values()), 0.0)
        for g, _ in GROUPS:
            v, src = (focus[g], "Given (QTOD)") if g in focus else (rest * mix.get(g, 0.0), how)
            if v:
                rows.append((year, code, r["Partner"], r["Branch"], r["Region"], r["Partner type"], GROUP_ACTIVITY[g], g,
                             v, "Gross", src))
    return pd.DataFrame(rows, columns=["Year", "Partner Code", "Partner", "Branch", "Region", "Partner type", "Activity",
                                       "Product group", "Target 100%", "Basis", "Source"])


def ssp_targets(master, year, issues):
    s = master.get("SSA Targets")
    if s is None:
        return pd.DataFrame()
    s = s.copy()
    s["SSP (as given)"] = s["SSP"].astype(str).str.strip()
    s["SSP"] = s["SSP (as given)"].replace(SSP_ALIASES)
    ch = s[s["SSP"] != s["SSP (as given)"]]
    if len(ch):
        issues.add("SSA Targets", "SSP names spelt differently from the sales data",
                   "renamed: " + ", ".join(sorted({f"{a} -> {b}" for a, b in zip(ch['SSP (as given)'], ch['SSP'])})))
    s["Branch (as given)"] = s["Branch"]
    s["Branch"] = s["Branch"].map(canon_branch)
    if "old Target" in s and s["old Target"].notna().any():
        issues.add("SSA Targets", "'old Target' column", "kept as 'Previous target' for reference")
    s = s.rename(columns={"old Target": "Previous target"})
    s["Basis"] = "Net"
    return s[["Year", "SSP", "Branch", "Target", "Previous target", "Basis", "SSP (as given)", "Branch (as given)"]]


def region_targets(master, issues):
    out = []
    z = master.get("Zone Targets")
    if z is not None:
        out.append(z.rename(columns={"Zone": "Region"}).assign(**{"Product group": "All"}))
    p = master.get("Product Targets")
    if p is not None:
        p = p.iloc[:, :4].copy()
        p.columns = ["Product group", "Region", "Year", "Target"]
        d = p[p.duplicated(["Product group", "Region", "Year"], keep=False)]
        if len(d):
            issues.add("Product Targets", "'East' appears twice per family (East1 and East2 both written as East)",
                       "added together as East", f"{len(d)} rows")
        p = p.groupby(["Year", "Region", "Product group"], as_index=False)["Target"].sum()
        out.append(p)
        if master["Product Targets"].shape[1] > 4:
            issues.add("Product Targets", "a list of family names to the right of the data", "not read")
    if not out:
        return pd.DataFrame()
    r = pd.concat(out, ignore_index=True)[["Year", "Region", "Product group", "Target"]]
    r["Basis"] = "Net"
    return r


def organize(master, qtod, sales, year, hs=None, partner_info=None):
    """master / qtod / hs: {sheet: DataFrame} of the target files. Returns ({sheet: DataFrame}, issues DataFrame)."""
    issues = Issues()
    sales = prepare_sales(sales)
    partner_info = partner_info or {}
    yearly_r, monthly_r, yearly_given, monthly_given, region = retail_targets(master, sales, year, issues)
    yearly_p = project_targets(hs, sales, year, issues)
    com = commits(hs, year, issues)
    if len(com) and len(yearly_p):
        # the commit sheet writes some Project branches in capitals (MUM, PUNE); use the target sheet's spelling
        names = {b.upper(): b for b in yearly_p["Branch"].unique()}
        p_rows = com["GTM"] == "Project"
        fixed = com.loc[p_rows, "Branch"].map(lambda b: names.get(str(b).upper(), b))
        changed = sorted(set(com.loc[p_rows, "Branch"][fixed != com.loc[p_rows, "Branch"]]))
        if changed:
            issues.add("Activity wise Month Target", "Project branch names in capitals: " + ", ".join(changed),
                       "matched to the Project targets' spelling")
        com.loc[p_rows, "Branch"] = fixed
    monthly_p = project_monthly(com, sales, year, issues, yearly_p)
    yearly = pd.concat([yearly_r, yearly_p], ignore_index=True)
    monthly = pd.concat([monthly_r, monthly_p], ignore_index=True)
    yearly, monthly = yearly[yearly["Target"].round(0) != 0], monthly[monthly["Target"].round(0) != 0]
    pt, q = partner_targets(qtod, master, sales, year, partner_info, region, issues)
    pps = partner_product_split(q, sales, year)
    branches = (pd.concat([yearly[["Branch", "Region", "GTM"]], monthly[["Branch", "Region", "GTM"]]])
                .drop_duplicates(["Branch", "GTM"]).sort_values(["GTM", "Region", "Branch"]))
    branches["Also written as"] = branches["Branch"].map(
        lambda b: ", ".join(sorted(k for k, v in BRANCH_ALIASES.items() if v == b)))
    tables = {"Branches": branches, "Yearly Targets": yearly, "Monthly Targets": monthly,
              "Partner Targets": pt, "Partner Product Split": pps, "SSP Targets": ssp_targets(master, year, issues),
              "Region Targets": region_targets(master, issues)}
    for k in tables:
        for c in ("Target", "Target 100%", "Target 108%"):
            if c in tables[k]:
                tables[k][c] = tables[k][c].astype(float).round(0)
    given = pd.DataFrame({"Yearly target (given)": yearly_given, "Monthly targets (given, all months)": monthly_given})
    return tables, pd.DataFrame(issues), given


# ====================================================================== workbook
def write_targets(path, tables, issues, given, year, base_year):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    green = "1F7A4D"
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    notes = [
        (f"Channel Targets {year}", "Retail, HS (Retail + Project) and Wiser targets in one place, one row per target."),
        ("", ""),
        ("Source", "Built from 'Master Targets File', 'QTOD and ATOD file' and the HS targets file. Every target is kept "
                   f"as given and only split over product groups by {base_year} sales; column 'Source' says which."),
        ("Amounts", "Rupees. 'Basis' says Net or Gross: yearly, monthly and SSP targets are Net (gross = net / 0.83); "
                    "partner targets are Gross."),
        ("Yearly Targets", "Year x branch x GTM x product group (activity alongside). Retail: 'Branch Targets'; Project: "
                           "the HS yearly targets."),
        ("Monthly Targets", "Month x branch x GTM x product group: the monthly commit. Retail: 'Monthly Targets' (Zeta and "
                            "Wiser as given); Project: the HS commits. Separate from the yearly target — the months "
                            "need not add up to it, and months not set yet are empty."),
        ("Partner Targets", "Partner x quarter x family ('All' = the partner's total, as in QTOD)."),
        ("Partner Product Split", "Each partner's yearly target split over product groups by its own "
                                  f"{base_year} mix — for planning, not a given target."),
        ("SSP Targets", "SSP x branch x year, names as in the sales data."),
        ("Region Targets", "Zone and product targets as given (region level)."),
        ("Issues", "Everything found in the source files and what was done about it."),
        ("Check", "Each sheet against the source file; the 'Difference' columns should be 0."),
        ("Editing", "Add rows at the bottom of a table; keep one row per key; use the dropdowns."),
    ]
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 120
    for r, (a, b) in enumerate(notes, 1):
        ws.cell(r, 1, a).font = Font(bold=True, size=14 if r == 1 else 11, color=green if r == 1 else "000000")
        ws.cell(r, 2, b).alignment = Alignment(wrap_text=True, vertical="top")

    fmt = {"Month": "mmm-yy", "Target": "#,##0", "Target 100%": "#,##0", "Target 108%": "#,##0", "Commit": "#,##0",
           "OB Commit": "#,##0", "Previous target": "#,##0"}

    def sheet(name, df, widths=None):
        w = wb.create_sheet(name)
        df = df.copy()
        if "Month" in df:
            df["Month"] = pd.to_datetime(df["Month"]).dt.date
        cols = list(df.columns)
        for j, c in enumerate(cols, 1):
            cell = w.cell(1, j, c)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=green)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            w.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, max(10, min(40, len(c) + 4)))
        for row in df.itertuples(index=False):
            w.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
        w.row_dimensions[1].height = 32
        w.freeze_panes = "A2"
        n = max(2, len(df) + 1)
        t = Table(displayName=re.sub(r"\W", "", name), ref=f"A1:{get_column_letter(len(cols))}{n}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        w.add_table(t)
        for j, c in enumerate(cols, 1):
            if c in fmt:
                for r in range(2, n + 1):
                    w.cell(r, j).number_format = fmt[c]
        return w, cols

    lists = {"GTM": ["Retail", "Project"], "Activity": ACTIVITIES, "Product group": [g for g, _ in GROUPS],
             "Family": ["All", "Miluz Zeta", "Wiser"], "Basis": ["Net", "Gross"], "Quarter": list(QUARTERS)}

    def dropdowns(w, cols):
        for c, vals in lists.items():
            if c in cols:
                letter = get_column_letter(cols.index(c) + 1)
                dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
                dv.add(f"{letter}2:{letter}50000")
                w.add_data_validation(dv)

    order = ["Branches", "Yearly Targets", "Monthly Targets", "Partner Targets", "Partner Product Split",
             "SSP Targets", "Region Targets"]
    for name in order:
        df = tables.get(name)
        if df is None or df.empty:
            continue
        w, cols = sheet(name, df, {"Source": 46, "Partner": 30, "Problem": 60})
        dropdowns(w, cols)
    if len(issues):
        w, _ = sheet("Issues", issues, {"Source sheet": 24, "Problem": 70, "What was done": 60, "Rows / values": 80})
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    # check: per Retail branch, each sheet against the source file
    yt, mt, pt = tables["Yearly Targets"], tables["Monthly Targets"], tables["Partner Targets"]
    chk = given.copy()
    chk["Yearly Targets sheet"] = yt[yt["GTM"] == "Retail"].groupby("Branch")["Target"].sum()
    chk["Difference (yearly)"] = (chk["Yearly Targets sheet"] - chk["Yearly target (given)"]).round(0)
    chk["Monthly Targets sheet"] = mt[mt["GTM"] == "Retail"].groupby("Branch")["Target"].sum()
    chk["Difference (monthly)"] = (chk["Monthly Targets sheet"] - chk["Monthly targets (given, all months)"]).round(0)
    chk["Months set"] = mt[mt["GTM"] == "Retail"].groupby("Branch")["Month"].nunique()
    chk["Months as % of yearly (info)"] = chk["Monthly Targets sheet"] / chk["Yearly target (given)"]
    chk["Partner targets 100% (gross)"] = pt[pt["Family"] == "All"].groupby("Branch")["Target 100%"].sum()
    chk = chk.reset_index().rename(columns={"index": "Branch"})
    w, cols = sheet("Check", chk, {"Branch": 14})
    for j, c in enumerate(cols, 1):
        f = "0%" if c.startswith("Months as") else ("0" if c == "Months set" else ("#,##0" if c != "Branch" else None))
        if f:
            for rr in range(2, len(chk) + 2):
                w.cell(rr, j).number_format = f
    wb.move_sheet("Check", offset=-(len(wb.sheetnames) - 2))
    if "Issues" in wb.sheetnames:
        wb.move_sheet("Issues", offset=-(wb.sheetnames.index("Issues") - 2))
    for w in wb.worksheets:
        w.sheet_properties.tabColor = green if w.title in ("Read me", "Check", "Issues") else "8FBC9F"
    wb.save(path)


def write_contribution(path, tables):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in tables.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("master")
    ap.add_argument("qtod")
    ap.add_argument("sales")
    ap.add_argument("out_dir")
    ap.add_argument("--hs")
    ap.add_argument("--year", type=int, default=dt.date.today().year)
    a = ap.parse_args()
    master = pd.read_excel(a.master, sheet_name=None)
    qtod = pd.read_excel(a.qtod, sheet_name=None)
    qtod = next(iter(qtod.values()))
    hs = pd.read_excel(a.hs, sheet_name=None) if a.hs else None
    sales = pd.read_parquet(a.sales) if a.sales.endswith(".parquet") else pd.read_pickle(a.sales)
    tables, issues, given = organize(master, qtod, sales, a.year, hs)
    write_targets(f"{a.out_dir}/Channel_Targets_{a.year}.xlsx", tables, issues, given, a.year, a.year - 1)
    write_contribution(f"{a.out_dir}/Sales_Contribution_{a.year - 1}.xlsx", contribution(prepare_sales(sales), a.year - 1))
    print(f"{len(issues)} issues; sheets: " + ", ".join(f"{k} {len(v)}" for k, v in tables.items()))


if __name__ == "__main__":
    main()
